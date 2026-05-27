# api/download.py
# GET /api/download?job_id=...
# Returns the generated workbook for the job.
#
# Mock: returns a tiny in-memory placeholder .xlsx (a minimal valid OOXML
# workbook authored with openpyxl) with the correct Content-Disposition naming
# the file {YYMMDD}_Ratings_{DealName}.xlsm and the macro-enabled MIME type
# application/vnd.ms-excel.sheet.macroEnabled.12.

from __future__ import annotations

import io
from http.server import BaseHTTPRequestHandler
from urllib.parse import parse_qs, urlparse

from openpyxl import Workbook

from _state import get_job


def _build_placeholder_workbook(deal_name: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ratings"
    ws["A1"] = "Arax Properties — Ratings"
    ws["A2"] = f"Deal: {deal_name}"
    ws["A4"] = "(Mock placeholder workbook. Real output preserves master macros via openpyxl keep_vba=True.)"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class handler(BaseHTTPRequestHandler):
    def do_GET(self):  # noqa: N802
        qs = parse_qs(urlparse(self.path).query)
        job_id = (qs.get("job_id") or [""])[0]
        if not job_id:
            self.send_response(400)
            self.send_header("Content-Type", "text/plain")
            self.end_headers()
            self.wfile.write(b"Missing job_id")
            return

        job = get_job(job_id)
        if job is None:
            self.send_response(404)
            self.send_header("Content-Type", "text/plain")
            self.end_headers()
            self.wfile.write(b"Job not found")
            return

        filename: str = job["workbook_filename"]
        deal_name: str = job["deal_name"]
        body = _build_placeholder_workbook(deal_name)

        self.send_response(200)
        self.send_header(
            "Content-Type",
            "application/vnd.ms-excel.sheet.macroEnabled.12",
        )
        self.send_header(
            "Content-Disposition",
            f'attachment; filename="{filename}"',
        )
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)
